"""
Lambda Combiner: 結果統合とExcel生成
S3から全Processor結果を収集し、Excelファイルに統合してS3にアップロード
"""

import os
import io
import json
import logging
import boto3
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Border, Alignment, Side
from openpyxl.utils import get_column_letter

# Configure logging
logger = logging.getLogger()
logger.setLevel(logging.INFO)

# AWS Clients
s3_client = boto3.client('s3')

# Environment variables
RESULTS_BUCKET = os.getenv('RESULTS_BUCKET', 'askul-sandbox-01-regulation-test-bucket')


def lambda_handler(event, context):
    """
    Lambda Combinerのメインハンドラー
    
    Expected event structure from Step Functions:
    {
        "execution_id": "20260204-143025-abc123",
        "input_bucket": "bucket-name",
        "input_key": "input/file.xlsx",
        "output_bucket": "bucket-name",
        "output_key": "output/file_checked.xlsx",
        "total_rows": 130,
        "results": [
            {"row_index": 0, "status": "success", "s3_key": "results/.../0.json"},
            {"row_index": 1, "status": "success", "s3_key": "results/.../1.json"},
            ...
        ]
    }
    
    Returns:
    {
        "statusCode": 200,
        "output_bucket": "bucket-name",
        "output_key": "output/file_checked.xlsx",
        "processed_rows": 130,
        "success_count": 128,
        "error_count": 2
    }
    """
    try:
        logger.info(f"Lambda Combiner started")
        
        execution_id = event.get('execution_id')
        input_bucket = event.get('input_bucket')
        input_key = event.get('input_key')
        output_bucket = event.get('output_bucket')
        output_key = event.get('output_key')
        total_rows = event.get('total_rows', 0)
        results = event.get('results', [])
        
        logger.info(f"Execution ID: {execution_id}")
        logger.info(f"Input: s3://{input_bucket}/{input_key}")
        logger.info(f"Output: s3://{output_bucket}/{output_key}")
        logger.info(f"Total rows: {total_rows}, Results received: {len(results)}")
        
        # Step 1: 元のExcelファイルをS3から読み込み
        logger.info("Downloading original Excel file from S3")
        excel_bytes = download_from_s3(input_bucket, input_key)
        
        # Step 2: Excelファイルを解析
        df, header_row, header_styles = parse_excel(excel_bytes)
        logger.info(f"Parsed Excel: {len(df)} rows, header at row {header_row}")
        
        # Step 3: S3から各行の処理結果を取得
        logger.info("Loading results from S3")
        row_results = load_results_from_s3(results)
        
        # 成功/エラーカウント
        success_count = sum(1 for r in results if r.get('status') == 'success')
        error_count = sum(1 for r in results if r.get('status') == 'error')
        
        logger.info(f"Success: {success_count}, Error: {error_count}")
        
        # Step 4: 結果をExcelに統合
        logger.info("Combining results into Excel")
        excel_output = combine_results_to_excel(df, row_results, header_row, header_styles)
        
        # Step 5: S3にアップロード
        logger.info(f"Uploading result to S3: {output_key}")
        upload_to_s3(output_bucket, output_key, excel_output)
        
        logger.info("Combiner completed successfully")
        
        return {
            'statusCode': 200,
            'output_bucket': output_bucket,
            'output_key': output_key,
            'processed_rows': total_rows,
            'success_count': success_count,
            'error_count': error_count
        }
        
    except Exception as e:
        logger.error(f"Combiner failed: {str(e)}", exc_info=True)
        raise


def download_from_s3(bucket, key):
    """S3からファイルをダウンロード"""
    try:
        response = s3_client.get_object(Bucket=bucket, Key=key)
        return response['Body'].read()
    except Exception as e:
        logger.error(f"Failed to download s3://{bucket}/{key}: {str(e)}")
        raise


def upload_to_s3(bucket, key, data):
    """S3にファイルをアップロード"""
    try:
        s3_client.put_object(
            Bucket=bucket,
            Key=key,
            Body=data,
            ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        logger.info(f"Uploaded to s3://{bucket}/{key}")
    except Exception as e:
        logger.error(f"Failed to upload to s3://{bucket}/{key}: {str(e)}")
        raise


def parse_excel(excel_bytes):
    """
    Excelファイルを解析
    
    Returns:
        tuple: (df, header_row, header_styles)
    """
    # openpyxlでワークブックを読み込み（.xlsm対応）
    wb = load_workbook(io.BytesIO(excel_bytes), keep_vba=True, data_only=True)
    ws = wb.active
    
    # ヘッダー行を検索（最大20行まで拡張）
    header_row = None
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
        if any('商品名' in str(cell) for cell in row if cell):
            header_row = row_idx
            break
    
    if header_row is None:
        raise ValueError("Header row with '商品名' not found in first 20 rows")
    
    # ヘッダーのスタイル情報を取得
    header_styles = {}
    for cell in ws[header_row]:
        if cell.value:
            header_styles[str(cell.value)] = {
                'fill': cell.fill.copy() if cell.fill else None,
                'font': cell.font.copy() if cell.font else None,
                'border': cell.border.copy() if cell.border else None,
                'alignment': cell.alignment.copy() if cell.alignment else None
            }
    
    # pandasでデータフレームとして読み込み
    df = pd.read_excel(io.BytesIO(excel_bytes), header=header_row-1, engine='openpyxl')
    
    return df, header_row, header_styles


def load_results_from_s3(results):
    """
    S3から各行の処理結果を読み込み
    
    Args:
        results: Step Functionsから渡されたresults配列
        
    Returns:
        list: 行ごとの結果辞書のリスト
    """
    row_results = []
    
    for result_info in results:
        row_index = result_info.get('row_index')
        s3_key = result_info.get('s3_key')
        status = result_info.get('status')
        
        if status == 'error':
            # エラーの場合
            error_msg = result_info.get('error', 'Unknown error')
            row_results.append({
                'row_index': row_index,
                'error': error_msg
            })
            logger.warning(f"Row {row_index}: Error - {error_msg}")
            continue
        
        # S3から結果JSONを取得
        try:
            obj = s3_client.get_object(Bucket=RESULTS_BUCKET, Key=s3_key)
            result_data = json.loads(obj['Body'].read())
            row_results.append(result_data)
            logger.debug(f"Row {row_index}: Loaded result from {s3_key}")
        except Exception as e:
            logger.error(f"Failed to load result from {s3_key}: {str(e)}")
            row_results.append({
                'row_index': row_index,
                'error': f"Failed to load result: {str(e)}"
            })
    
    # row_indexでソート
    row_results.sort(key=lambda x: x.get('row_index', 0))
    
    return row_results


def combine_results_to_excel(df, row_results, header_row, header_styles):
    """
    結果をExcelに統合
    
    Args:
        df: 元のDataFrame
        row_results: 行ごとの結果辞書のリスト
        header_row: ヘッダー行番号
        header_styles: ヘッダーのスタイル情報
        
    Returns:
        bytes: Excelファイルのバイト列
    """
    # チェック対象列のリスト（SKILL.md 指定の8列）
    check_columns = [
        '変更後_キャッチコピーBtoB',
        '変更後_商品の特徴BtoB',
        '変更後_短いキャッチコピーBtoB',
        '変更後_MDおすすめコメントBtoB',
        '変更後_キャッチコピーBtoC',
        '変更後_商品の特徴BtoC',
        '変更後_MDおすすめコメントBtoC',
        '変更後_短いキャッチコピーBtoC'
    ]
    
    # 各項目のチェック結果を格納する辞書
    check_results = {col: [''] * len(df) for col in check_columns}
    
    # row_resultsから各項目の結果を抽出
    for result_data in row_results:
        row_index = result_data.get('row_index', -1)
        
        if row_index < 0 or row_index >= len(df):
            logger.warning(f"Invalid row_index: {row_index}")
            continue
        
        # エラーの場合
        if 'error' in result_data:
            error_msg = result_data['error']
            for col in check_columns:
                check_results[col][row_index] = f"エラー: {error_msg}"
            continue
        
        # 各列のチェック結果を取得
        for col in check_columns:
            result_col_name = f"{col}_チェック結果"
            if result_col_name in result_data:
                check_results[col][row_index] = result_data[result_col_name]
    
    # 新しい列順を構築：各チェック対象列の直後にチェック結果列を挿入
    # 元の列が存在しない場合でも結果列は追加
    all_columns = list(df.columns)
    new_df_dict = {}
    added_check_columns = set()  # 追加済みの結果列を追跡
    
    # まず元の列をループして、存在する列の直後に結果列を挿入
    for col in all_columns:
        # 元の列を追加
        new_df_dict[col] = df[col]
        # チェック対象列の場合、直後にチェック結果列を追加
        if col in check_columns:
            result_col_name = f'{col}_チェック結果'
            new_df_dict[result_col_name] = check_results[col]
            added_check_columns.add(col)
    
    # 元の列に存在しないチェック対象列の結果列も末尾に追加
    for col in check_columns:
        if col not in added_check_columns:
            result_col_name = f'{col}_チェック結果'
            new_df_dict[result_col_name] = check_results[col]
    
    # 新しいDataFrameを作成
    new_df = pd.DataFrame(new_df_dict)
    
    # ExcelWriterで書き出し（スタイル適用）
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        new_df.to_excel(writer, index=False, sheet_name='Sheet1', startrow=header_row-1)
        
        # ワークシートを取得
        ws = writer.sheets['Sheet1']
        
        # ヘッダー行のスタイルを適用
        apply_header_styles(ws, header_row, header_styles, check_columns)
        
        # 列幅を調整
        adjust_column_widths(ws)
    
    output.seek(0)
    return output.getvalue()


def apply_header_styles(ws, header_row, header_styles, check_columns):
    """ヘッダー行にスタイルを適用"""
    # デフォルトのヘッダースタイル
    default_fill = PatternFill(start_color='D3D3D3', end_color='D3D3D3', fill_type='solid')
    default_font = Font(bold=True)
    default_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    for cell in ws[header_row]:
        if cell.value:
            col_name = str(cell.value)
            
            # 元の列のスタイルを適用
            if col_name in header_styles:
                styles = header_styles[col_name]
                if styles.get('fill'):
                    cell.fill = styles['fill']
                else:
                    cell.fill = default_fill
                if styles.get('font'):
                    cell.font = styles['font']
                else:
                    cell.font = default_font
                if styles.get('border'):
                    cell.border = styles['border']
                else:
                    cell.border = thin_border
                if styles.get('alignment'):
                    cell.alignment = styles['alignment']
                else:
                    cell.alignment = default_alignment
            
            # チェック結果列の場合、異なる色で塗る
            elif col_name.endswith('_チェック結果'):
                cell.fill = PatternFill(start_color='FFE6CC', end_color='FFE6CC', fill_type='solid')
                cell.font = default_font
                cell.border = thin_border
                cell.alignment = default_alignment
            
            # その他の列
            else:
                cell.fill = default_fill
                cell.font = default_font
                cell.border = thin_border
                cell.alignment = default_alignment


def adjust_column_widths(ws):
    """列幅を調整"""
    for col_idx, column in enumerate(ws.columns, start=1):
        col_letter = get_column_letter(col_idx)
        max_length = 0
        
        for cell in column:
            if cell.value:
                # セル値の長さを計算
                cell_value = str(cell.value)
                # 改行で分割して最大長を取得
                lines = cell_value.split('\n')
                max_line_length = max(len(line) for line in lines)
                max_length = max(max_length, max_line_length)
        
        # 列幅を設定（最小10、最大50）
        adjusted_width = min(max(max_length + 2, 10), 50)
        ws.column_dimensions[col_letter].width = adjusted_width
